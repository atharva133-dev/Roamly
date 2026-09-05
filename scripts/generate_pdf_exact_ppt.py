import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_pdf_exact_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette matching the exact PDF template
    BG_COLOR = RGBColor(233, 239, 248)       # #E9EFF8 Ice Blue Background
    MAIN_TITLE_COLOR = RGBColor(44, 62, 107) # #2C3E6B Navy Dark Slate Blue
    WHITE_CARD = RGBColor(255, 255, 255)     # #FFFFFF White
    GRAY_BLUE_BOX = RGBColor(204, 215, 230)  # #CCD7E6 Grayish Blue Content Container
    TEXT_DARK = RGBColor(30, 41, 59)         # #1E293B Slate-800
    TEXT_MUTED = RGBColor(71, 85, 105)       # #475569 Slate-600
    ACCENT_INDIGO = RGBColor(79, 70, 229)   # #4F46E5 Indigo
    ACCENT_BLUE = RGBColor(37, 99, 235)     # #2563EB Blue

    blank_layout = prs.slide_layouts[6]

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_slide_header(slide, title_text):
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.9))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = MAIN_TITLE_COLOR
        p.font.name = "Segoe UI"

    # =========================================================================
    # PAGE 1: HackCelestial 3.0 Title & Abstract Slide
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_background(slide1)

    # University Header Text
    tx_uni = slide1.shapes.add_textbox(Inches(0.8), Inches(0.3), Inches(11.733), Inches(0.5))
    tf_uni = tx_uni.text_frame
    tf_uni.word_wrap = True
    p_u1 = tf_uni.paragraphs[0]
    p_u1.text = "Mahatma Education Society's"
    p_u1.font.size = Pt(14)
    p_u1.font.bold = True
    p_u1.font.color.rgb = MAIN_TITLE_COLOR
    p_u1.alignment = PP_ALIGN.CENTER
    p_u1.font.name = "Segoe UI"

    p_u2 = tf_uni.add_paragraph()
    p_u2.text = "PILLAI UNIVERSITY"
    p_u2.font.size = Pt(26)
    p_u2.font.bold = True
    p_u2.font.color.rgb = MAIN_TITLE_COLOR
    p_u2.alignment = PP_ALIGN.CENTER
    p_u2.font.name = "Segoe UI"

    # HackCelestial 3.0 Big Title
    tx_hack = slide1.shapes.add_textbox(Inches(0.8), Inches(1.15), Inches(11.733), Inches(0.8))
    tf_hack = tx_hack.text_frame
    tf_hack.word_wrap = True
    p_h = tf_hack.paragraphs[0]
    p_h.text = "HackCelestial 3.0"
    p_h.font.size = Pt(44)
    p_h.font.bold = True
    p_h.font.color.rgb = MAIN_TITLE_COLOR
    p_h.alignment = PP_ALIGN.CENTER
    p_h.font.name = "Segoe UI"

    # Left Top Card: Team Name
    c1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(5.6), Inches(1.8))
    c1.fill.solid()
    c1.fill.fore_color.rgb = WHITE_CARD
    c1.line.fill.background()

    tx_team = slide1.shapes.add_textbox(Inches(1.1), Inches(2.5), Inches(5.0), Inches(1.2))
    tf_team = tx_team.text_frame
    tf_team.word_wrap = True
    p_t1 = tf_team.paragraphs[0]
    p_t1.text = "Team Name:"
    p_t1.font.size = Pt(22)
    p_t1.font.bold = True
    p_t1.font.color.rgb = MAIN_TITLE_COLOR
    p_t1.font.name = "Segoe UI"

    p_t2 = tf_team.add_paragraph()
    p_t2.text = "Ghayal Ghoda"
    p_t2.font.size = Pt(28)
    p_t2.font.bold = True
    p_t2.font.color.rgb = TEXT_DARK
    p_t2.font.name = "Segoe UI"

    # Left Bottom Card: Problem Statement Title
    c2 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.7))
    c2.fill.solid()
    c2.fill.fore_color.rgb = WHITE_CARD
    c2.line.fill.background()

    tx_ps = slide1.shapes.add_textbox(Inches(1.1), Inches(4.5), Inches(5.0), Inches(2.3))
    tf_ps = tx_ps.text_frame
    tf_ps.word_wrap = True
    p_ps1 = tf_ps.paragraphs[0]
    p_ps1.text = "Problem Statement Title:"
    p_ps1.font.size = Pt(20)
    p_ps1.font.bold = True
    p_ps1.font.color.rgb = MAIN_TITLE_COLOR
    p_ps1.font.name = "Segoe UI"
    p_ps1.space_after = Pt(6)

    p_ps2 = tf_ps.add_paragraph()
    p_ps2.text = "Personalized Dynamic Tour Planning & Tour Operations Platform - PS-07"
    p_ps2.font.size = Pt(18)
    p_ps2.font.bold = True
    p_ps2.font.color.rgb = TEXT_DARK
    p_ps2.font.name = "Segoe UI"

    # Right Card: Pitch Summary & Abstract
    c3 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.7), Inches(2.2), Inches(5.833), Inches(4.8))
    c3.fill.solid()
    c3.fill.fore_color.rgb = WHITE_CARD
    c3.line.fill.background()

    tx_abs = slide1.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.233), Inches(4.4))
    tf_abs = tx_abs.text_frame
    tf_abs.word_wrap = True
    p_a1 = tf_abs.paragraphs[0]
    p_a1.text = "03. Pitch Summary"
    p_a1.font.size = Pt(14)
    p_a1.font.bold = True
    p_a1.font.color.rgb = ACCENT_INDIGO
    p_a1.font.name = "Segoe UI"

    p_a2 = tf_abs.add_paragraph()
    p_a2.text = "Abstract:"
    p_a2.font.size = Pt(24)
    p_a2.font.bold = True
    p_a2.font.color.rgb = MAIN_TITLE_COLOR
    p_a2.font.name = "Segoe UI"
    p_a2.space_after = Pt(4)

    abstract_bullets = [
        "Tripzzy is a personalized dynamic tour planning and tour operations platform.",
        "It helps users create customized travel itineraries based on their preferences (budget, style, dates, interests).",
        "The platform simplifies trip planning, scheduling, multi-city routing, and tour operations management.",
        "Integrates a high-performance Redis + BullMQ asynchronous queue system to eliminate AI latency during peak loads.",
        "It provides a dynamic and user-friendly interface with interactive 3D Canvas Globe, Map view, and Calendar schedule.",
        "Users can efficiently manage destinations, activities, localized INR (₹) budgets, and overall trip schedules.",
        "The solution aims to reduce the time and complexity involved in planning tours.",
        "It offers a centralized platform for a smooth, transparent, and personalized travel experience.",
        "Tripzzy is designed to make tour planning flexible, convenient, efficient, and scalable for users and tour operators."
    ]

    for bullet in abstract_bullets:
        p_b = tf_abs.add_paragraph()
        p_b.text = "•  " + bullet
        p_b.font.size = Pt(11)
        p_b.font.color.rgb = TEXT_DARK
        p_b.font.name = "Segoe UI"
        p_b.space_after = Pt(2)

    # =========================================================================
    # PAGE 2: Proposed Solution
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_background(slide2)
    add_slide_header(slide2, "Proposed Solution")

    box2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.6))
    box2.fill.solid()
    box2.fill.fore_color.rgb = GRAY_BLUE_BOX
    box2.line.fill.background()

    tx_s2 = slide2.shapes.add_textbox(Inches(1.2), Inches(1.7), Inches(10.933), Inches(5.0))
    tf_s2 = tx_s2.text_frame
    tf_s2.word_wrap = True

    p_sol_title = tf_s2.paragraphs[0]
    p_sol_title.text = "Detailed explanation of the proposed solution:"
    p_sol_title.font.size = Pt(18)
    p_sol_title.font.bold = True
    p_sol_title.font.color.rgb = MAIN_TITLE_COLOR
    p_sol_title.font.name = "Segoe UI"
    p_sol_title.space_after = Pt(10)

    solution_bullets = [
        "Tripzzy is a smart, personalized tour planning platform designed to simplify the entire travel-planning process.",
        "Users can enter their destination, budget, travel dates, interests, accommodation, transportation, and special preferences.",
        "The system generates a personalized itinerary with suitable places, morning/afternoon/evening activities, and schedules using Google Gemini LLM.",
        "Integrates a Redis + BullMQ asynchronous task queue with multi-tier Dead Letter Queues (DLQ) to process heavy AI requests without lag.",
        "It enables users to organize and modify their trip plan dynamically according to changing requirements.",
        "The platform provides a centralized view of the complete tour (3D Canvas Globe, Map view, Calendar timeline), reducing the need to use multiple applications.",
        "Automatically standardizes all financial costs into localized Indian Rupees (INR ₹) using locale-aware formatting.",
        "Tripzzy focuses on making travel planning faster, flexible, convenient, and user-friendly.",
        "The solution can also support tour operators in managing, coordinating, and analyzing tour activities via an admin dashboard.",
        "Overall, Tripzzy aims to provide a seamless, customized, highly efficient, and crash-proof travel-planning experience for users."
    ]

    for pt in solution_bullets:
        p = tf_s2.add_paragraph()
        p.text = "•  " + pt
        p.font.size = Pt(14)
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"
        p.space_after = Pt(8)

    # =========================================================================
    # PAGE 3: Flow Chart / Architecture
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_background(slide3)
    add_slide_header(slide3, "Flow Chart / Architecture")

    box3 = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.6))
    box3.fill.solid()
    box3.fill.fore_color.rgb = GRAY_BLUE_BOX
    box3.line.fill.background()

    tx_s3 = slide3.shapes.add_textbox(Inches(1.1), Inches(1.6), Inches(11.133), Inches(5.2))
    tf_s3 = tx_s3.text_frame
    tf_s3.word_wrap = True

    p_f1 = tf_s3.paragraphs[0]
    p_f1.text = "Flow Chart & Architecture: End-to-End System Data Flow"
    p_f1.font.size = Pt(18)
    p_f1.font.bold = True
    p_f1.font.color.rgb = MAIN_TITLE_COLOR
    p_f1.font.name = "Segoe UI"
    p_f1.space_after = Pt(8)

    # Visual Flowchart Boxes
    flow_steps = [
        ("1. USER INPUT", "Destinations, Dates, Budget, Style", MAIN_TITLE_COLOR),
        ("2. NEXT.JS API", "Clerk Auth & Payload Validation", MAIN_TITLE_COLOR),
        ("3. REDIS & BULLMQ", "SHA-256 Hash Cache & Queue", MAIN_TITLE_COLOR),
        ("4. GEMINI AI", "LLM Generation & INR Converter", MAIN_TITLE_COLOR),
        ("5. POSTGRES DB", "Prisma ORM Persistence", MAIN_TITLE_COLOR),
        ("6. CLIENT RENDER", "3D Globe, Map & Calendar", MAIN_TITLE_COLOR)
    ]

    f_lefts = [Inches(1.1), Inches(4.8), Inches(8.5)]
    f_tops = [Inches(2.2), Inches(3.6)]

    for idx, (step_title, step_desc, box_col) in enumerate(flow_steps):
        s_left = f_lefts[idx % 3]
        s_top = f_tops[idx // 3]

        # Draw card for each step inside slide 3
        scard = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, s_left, s_top, Inches(3.4), Inches(1.2))
        scard.fill.solid()
        scard.fill.fore_color.rgb = WHITE_CARD
        scard.line.color.rgb = MAIN_TITLE_COLOR
        scard.line.width = Pt(1.5)

        stx = slide3.shapes.add_textbox(s_left + Inches(0.1), s_top + Inches(0.1), Inches(3.2), Inches(1.0))
        stf = stx.text_frame
        stf.word_wrap = True

        sp1 = stf.paragraphs[0]
        sp1.text = step_title
        sp1.font.size = Pt(13)
        sp1.font.bold = True
        sp1.font.color.rgb = ACCENT_BLUE
        sp1.font.name = "Segoe UI"

        sp2 = stf.add_paragraph()
        sp2.text = step_desc
        sp2.font.size = Pt(11)
        sp2.font.color.rgb = TEXT_DARK
        sp2.font.name = "Segoe UI"

    # Architecture Explanation Bullets
    tx_arch = slide3.shapes.add_textbox(Inches(1.1), Inches(5.0), Inches(11.133), Inches(1.8))
    tf_arch = tx_arch.text_frame
    tf_arch.word_wrap = True

    arch_bullets = [
        "Client Layer: Next.js 15 App Router with React 19, TypeScript, Tailwind CSS, Shadcn UI & Framer Motion.",
        "Async Queue Layer: Redis key-value cache + BullMQ workers (Main Queue + DLQ1 + DLQ2 retries) for zero-latency execution.",
        "Database & Auth Layer: PostgreSQL database managed via Prisma ORM connected with Clerk Authentication.",
        "AI & Currency Engine: Google Gemini LLM API integrated with fallback handling and automated INR (₹) currency conversion."
    ]

    for pt in arch_bullets:
        p = tf_arch.add_paragraph()
        p.text = "•  " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"
        p.space_after = Pt(3)

    # =========================================================================
    # PAGE 4: Innovation and Unique Functionality
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_background(slide4)
    add_slide_header(slide4, "Innovation and Unique Functionality")

    box4 = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.6))
    box4.fill.solid()
    box4.fill.fore_color.rgb = GRAY_BLUE_BOX
    box4.line.fill.background()

    tx_s4 = slide4.shapes.add_textbox(Inches(1.2), Inches(1.7), Inches(10.933), Inches(5.0))
    tf_s4 = tx_s4.text_frame
    tf_s4.word_wrap = True

    p_in1 = tf_s4.paragraphs[0]
    p_in1.text = "Innovative Aspects:"
    p_in1.font.size = Pt(18)
    p_in1.font.bold = True
    p_in1.font.color.rgb = MAIN_TITLE_COLOR
    p_in1.font.name = "Segoe UI"
    p_in1.space_after = Pt(4)

    inn_bullets = [
        "Crash-Proof 3D Canvas Globe Engine: Custom-built 2D Canvas rendering engine with 3D matrix mathematics, specular shading, animated flight path arc vectors, and pulsing markers—completely immune to WebGL context crashes across low-end mobile devices.",
        "Resilient Asynchronous Task Queueing: Integrated Redis + BullMQ queue system with SHA-256 request hashing to return cached trip plans instantly and multi-tier Dead Letter Queues (DLQ1 & DLQ2) for fail-safe background processing.",
        "Automated INR Currency Normalizer: Built-in parser that automatically detects foreign currencies ($ USD, € EUR, £ GBP, ¥ JPY) in AI model outputs and converts all costs into Indian Rupees (INR ₹) using locale-aware formatting."
    ]

    for pt in inn_bullets:
        p = tf_s4.add_paragraph()
        p.text = "•  " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"
        p.space_after = Pt(6)

    p_uf1 = tf_s4.add_paragraph()
    p_uf1.text = "Unique Features:"
    p_uf1.font.size = Pt(18)
    p_uf1.font.bold = True
    p_uf1.font.color.rgb = MAIN_TITLE_COLOR
    p_uf1.font.name = "Segoe UI"
    p_uf1.space_after = Pt(4)

    uniq_bullets = [
        "Synchronized Dual Map & Calendar View: Real-time synchronization between spatial geographical locations (Map view) and temporal daily schedules (Calendar view).",
        "Community Trip Inspiration Hub: Public trip discovery network allowing users to search, save, and draw inspiration from itineraries generated by other travelers.",
        "Unified Tour Operations & Admin Analytics Dashboard: Centralized portal providing tour operators with real-time analytics on trending cities, activity distributions, user metrics, and reviews."
    ]

    for pt in uniq_bullets:
        p = tf_s4.add_paragraph()
        p.text = "•  " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"
        p.space_after = Pt(6)

    # =========================================================================
    # PAGE 5: Technical Details
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_background(slide5)
    add_slide_header(slide5, "Technical Details")

    box5 = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.6))
    box5.fill.solid()
    box5.fill.fore_color.rgb = GRAY_BLUE_BOX
    box5.line.fill.background()

    tx_s5 = slide5.shapes.add_textbox(Inches(1.2), Inches(1.7), Inches(10.933), Inches(5.0))
    tf_s5 = tx_s5.text_frame
    tf_s5.word_wrap = True

    p_t1 = tf_s5.paragraphs[0]
    p_t1.text = "Frameworks & Technologies:"
    p_t1.font.size = Pt(18)
    p_t1.font.bold = True
    p_t1.font.color.rgb = MAIN_TITLE_COLOR
    p_t1.font.name = "Segoe UI"
    p_t1.space_after = Pt(2)

    tech_bullets = [
        "Frontend: Next.js 15 (App Router), React 19, TypeScript, Tailwind CSS, Shadcn UI, Framer Motion, Lucide Icons.",
        "Backend & Async Queues: Node.js, Next.js Serverless Route Handlers, Redis (ioredis), BullMQ Worker Queues, Google Generative AI (Gemini SDK).",
        "Database & Authentication: PostgreSQL database, Prisma ORM (relational schema for users, trips, stops, activities, budgets, reviews), Clerk Auth."
    ]

    for pt in tech_bullets:
        p = tf_s5.add_paragraph()
        p.text = "•  " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"
        p.space_after = Pt(4)

    p_dep = tf_s5.add_paragraph()
    p_dep.text = "Deployment:"
    p_dep.font.size = Pt(18)
    p_dep.font.bold = True
    p_dep.font.color.rgb = MAIN_TITLE_COLOR
    p_dep.font.name = "Segoe UI"
    p_dep.space_after = Pt(2)

    dep_bullets = [
        "Vercel Cloud Platform (Serverless Next.js functions & edge network).",
        "Upstash Serverless Redis (Low-latency managed key-value store for caching & BullMQ queues).",
        "Neon Serverless PostgreSQL (Auto-scaling cloud relational database)."
    ]

    for pt in dep_bullets:
        p = tf_s5.add_paragraph()
        p.text = "•  " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"
        p.space_after = Pt(4)

    p_cost = tf_s5.add_paragraph()
    p_cost.text = "Cost:"
    p_cost.font.size = Pt(18)
    p_cost.font.bold = True
    p_cost.font.color.rgb = MAIN_TITLE_COLOR
    p_cost.font.name = "Segoe UI"
    p_cost.space_after = Pt(2)

    cost_bullets = [
        "Vercel Serverless: $0 - $20 / month | Upstash Redis: $0 - $10 / month | Neon PostgreSQL: $0 - $25 / month | Gemini AI API: Free Tier / Pay-as-you-go.",
        "Total Operational Cost: ~$0 – $55 / month (Highly cost-effective, serverless infrastructure with zero idle compute costs)."
    ]

    for pt in cost_bullets:
        p = tf_s5.add_paragraph()
        p.text = "•  " + pt
        p.font.size = Pt(13)
        p.font.color.rgb = TEXT_DARK
        p.font.name = "Segoe UI"
        p.space_after = Pt(4)

    # =========================================================================
    # PAGE 6: Existing Solutions and Comparison
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_background(slide6)
    add_slide_header(slide6, "Existing Solutions and Comparison")

    box6 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.6))
    box6.fill.solid()
    box6.fill.fore_color.rgb = GRAY_BLUE_BOX
    box6.line.fill.background()

    tx_s6 = slide6.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.333), Inches(0.5))
    tf_s6 = tx_s6.text_frame
    p_c1 = tf_s6.paragraphs[0]
    p_c1.text = "Existing Solutions and Comparison Matrix"
    p_c1.font.size = Pt(18)
    p_c1.font.bold = True
    p_c1.font.color.rgb = MAIN_TITLE_COLOR
    p_c1.font.name = "Segoe UI"

    # Comparison Table inside Box
    rows, cols = 7, 6
    left, top, width, height = Inches(1.0), Inches(2.2), Inches(11.333), Inches(4.5)
    table_shape = slide6.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(2.533)
    table.columns[1].width = Inches(1.9)
    table.columns[2].width = Inches(1.7)
    table.columns[3].width = Inches(1.7)
    table.columns[4].width = Inches(1.75)
    table.columns[5].width = Inches(1.75)

    headers = [
        "Feature / Dimension", 
        "Tripzzy (Roamly)", 
        "TripIt", 
        "Wanderlog", 
        "Generic ChatGPT", 
        "Tour Operators"
    ]

    matrix_data = [
        ["Real-time AI Itinerary Synthesis", "✅ Instant & Dynamic", "❌ Manual Forwarding", "⚠️ Limited AI", "⚠️ Unstructured Text", "❌ Slow / Fixed"],
        ["Asynchronous Queue (BullMQ)", "✅ Built-in Queue", "❌ None", "❌ None", "❌ None", "❌ None"],
        ["Crash-Proof 3D Globe Visualizer", "✅ Interactive 3D Canvas", "❌ Static Maps", "⚠️ Basic 2D Pin", "❌ Text Only", "❌ Printed Papers"],
        ["Localized Currency (INR ₹)", "✅ Automated Conversion", "❌ USD / Local Only", "⚠️ Mixed Currency", "⚠️ Manual Prompting", "⚠️ Fixed Quotes"],
        ["Multi-City Route Auto-Optimizer", "✅ Smart Logical Splits", "❌ Manual Layout", "⚠️ Semi-Automated", "⚠️ Inconsistent", "⚠️ Pre-packaged"],
        ["Integrated Admin & Community", "✅ Analytics + Sharing", "❌ Personal Only", "⚠️ Basic Sharing", "❌ None", "❌ Closed Systems"]
    ]

    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = MAIN_TITLE_COLOR if col_idx != 1 else ACCENT_BLUE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Segoe UI"

    for row_idx, row_content in enumerate(matrix_data):
        for col_idx, text in enumerate(row_content):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            if col_idx == 0:
                cell.fill.fore_color.rgb = RGBColor(241, 245, 249)
            elif col_idx == 1:
                cell.fill.fore_color.rgb = RGBColor(239, 246, 255)
            else:
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if row_idx % 2 == 0 else RGBColor(248, 250, 252)

            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(10)
            p.font.name = "Segoe UI"
            if col_idx == 0:
                p.font.bold = True
                p.font.color.rgb = TEXT_DARK
                p.alignment = PP_ALIGN.LEFT
            elif col_idx == 1:
                p.font.bold = True
                p.font.color.rgb = ACCENT_BLUE
                p.alignment = PP_ALIGN.CENTER
            else:
                p.font.color.rgb = TEXT_MUTED
                p.alignment = PP_ALIGN.CENTER

    output_filename = "Tripzzy_PS07_Ghayal_Ghoda_HackCelestial_Deck.pptx"
    output_path = os.path.join(os.getcwd(), output_filename)
    prs.save(output_path)
    print(f"Presentation saved successfully at: {output_path}")

    # Copy to artifacts directory if available
    artifact_dir = r"C:\Users\Atharva\.gemini\antigravity-ide\brain\00846d85-d373-421f-bf87-f6307f487954"
    if os.path.exists(artifact_dir):
        import shutil
        shutil.copy(output_path, os.path.join(artifact_dir, output_filename))
        print(f"Presentation copied to artifacts directory: {os.path.join(artifact_dir, output_filename)}")

if __name__ == "__main__":
    build_pdf_exact_presentation()
