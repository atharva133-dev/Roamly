import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color Palette
    BG_COLOR = RGBColor(241, 245, 249)      # #f1f5f9 slate light
    HEADER_BG = RGBColor(30, 41, 59)        # #1e293b dark slate
    PRIMARY_TEXT = RGBColor(15, 23, 42)      # #0f172a slate-900
    SECONDARY_TEXT = RGBColor(71, 85, 105)   # #475569 slate-600
    CARD_BG = RGBColor(255, 255, 255)       # white
    CARD_BORDER = RGBColor(203, 213, 225)   # slate-300
    ACCENT_BLUE = RGBColor(37, 99, 235)     # #2563eb blue-600
    ACCENT_INDIGO = RGBColor(79, 70, 229)   # #4f46e5 indigo-600
    ACCENT_GREEN = RGBColor(16, 185, 129)   # #10b981 emerald-500
    ACCENT_PURPLE = RGBColor(147, 51, 234)  # #9333ea purple-600

    blank_layout = prs.slide_layouts[6]

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="HackCelestial 3.0 • Pillai University"):
        # Top banner bar
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.1))
        banner.fill.solid()
        banner.fill.fore_color.rgb = HEADER_BG
        banner.line.fill.background()

        # Banner text
        txBox = slide.shapes.add_textbox(Inches(0.6), Inches(0.12), Inches(12.133), Inches(0.85))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0

        p1 = tf.paragraphs[0]
        p1.text = category_text.upper()
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_GREEN
        p1.font.name = "Calibri"

        p2 = tf.add_paragraph()
        p2.text = title_text
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = RGBColor(255, 255, 255)
        p2.font.name = "Calibri"

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.5)
        else:
            card.line.fill.background()
        return card

    # =========================================================================
    # SLIDE 1: Pitch Summary (Title & Abstract)
    # =========================================================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_background(slide1)

    # Top Header
    header_box = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.2))
    header_box.fill.solid()
    header_box.fill.fore_color.rgb = HEADER_BG
    header_box.line.fill.background()

    htx = slide1.shapes.add_textbox(Inches(0.6), Inches(0.15), Inches(12.133), Inches(0.9))
    htf = htx.text_frame
    htf.word_wrap = True
    hp1 = htf.paragraphs[0]
    hp1.text = "MAHATMA EDUCATION SOCIETY'S • PILLAI UNIVERSITY • TECH ALEGRIA"
    hp1.font.size = Pt(11)
    hp1.font.bold = True
    hp1.font.color.rgb = ACCENT_GREEN
    hp1.font.name = "Calibri"

    hp2 = htf.add_paragraph()
    hp2.text = "HackCelestial 3.0 — Official Pitch Deck"
    hp2.font.size = Pt(24)
    hp2.font.bold = True
    hp2.font.color.rgb = RGBColor(255, 255, 255)
    hp2.font.name = "Calibri"

    # Left Column: Team & Problem Statement
    add_card(slide1, Inches(0.6), Inches(1.5), Inches(5.5), Inches(2.5))
    tx1 = slide1.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(5.1), Inches(2.1))
    tf1 = tx1.text_frame
    tf1.word_wrap = True
    tp1 = tf1.paragraphs[0]
    tp1.text = "TEAM NAME"
    tp1.font.size = Pt(11)
    tp1.font.bold = True
    tp1.font.color.rgb = ACCENT_BLUE

    tp2 = tf1.add_paragraph()
    tp2.text = "Ghayal Ghoda"
    tp2.font.size = Pt(26)
    tp2.font.bold = True
    tp2.font.color.rgb = PRIMARY_TEXT

    add_card(slide1, Inches(0.6), Inches(4.3), Inches(5.5), Inches(2.7))
    tx2 = slide1.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(5.1), Inches(2.3))
    tf2 = tx2.text_frame
    tf2.word_wrap = True
    pp1 = tf2.paragraphs[0]
    pp1.text = "PROBLEM STATEMENT TITLE & ID"
    pp1.font.size = Pt(11)
    pp1.font.bold = True
    pp1.font.color.rgb = ACCENT_INDIGO

    pp2 = tf2.add_paragraph()
    pp2.text = "Personalized Dynamic Tour Planning & Tour Operations Platform — PS-07"
    pp2.font.size = Pt(18)
    pp2.font.bold = True
    pp2.font.color.rgb = PRIMARY_TEXT

    pp3 = tf2.add_paragraph()
    pp3.text = "\nProject Name: Tripzzy (Roamly)"
    pp3.font.size = Pt(14)
    pp3.font.bold = True
    pp3.font.color.rgb = ACCENT_BLUE

    # Right Column: Pitch Summary / Abstract
    add_card(slide1, Inches(6.4), Inches(1.5), Inches(6.333), Inches(5.5))
    tx3 = slide1.shapes.add_textbox(Inches(6.7), Inches(1.7), Inches(5.733), Inches(5.1))
    tf3 = tx3.text_frame
    tf3.word_wrap = True
    ap1 = tf3.paragraphs[0]
    ap1.text = "03. PITCH SUMMARY & ABSTRACT"
    ap1.font.size = Pt(11)
    ap1.font.bold = True
    ap1.font.color.rgb = ACCENT_GREEN

    ap2 = tf3.add_paragraph()
    ap2.text = "Abstract:"
    ap2.font.size = Pt(20)
    ap2.font.bold = True
    ap2.font.color.rgb = PRIMARY_TEXT

    abstract_points = [
        "Tripzzy (Roamly) is an end-to-end personalized dynamic tour planning and tour operations platform.",
        "It empowers users to instantly generate customized, multi-city travel itineraries based on budget, travel style, dates, accommodation, and specific interests.",
        "The platform simplifies complex trip planning, multi-city routing, scheduling, and tour operations management in a unified interface.",
        "Integrates a high-performance Redis + BullMQ asynchronous task queue to eliminate LLM processing latency and ensure zero-wait user experience.",
        "Features an interactive 3D Canvas Globe, synchronized Map & Calendar schedule views, and standardized INR (₹) cost calculations.",
        "Provides a centralized administrative dashboard for tour operators to monitor trending destinations, activity distribution, and user reviews.",
        "Tripzzy makes tour planning fast, flexible, convenient, transparent, and highly efficient for both travelers and tour operators."
    ]

    for pt in abstract_points:
        p = tf3.add_paragraph()
        p.text = "•  " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = SECONDARY_TEXT
        p.space_after = Pt(6)

    # =========================================================================
    # SLIDE 2: Proposed Solution
    # =========================================================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_background(slide2)
    add_header(slide2, "Proposed Solution — Tripzzy (Roamly)")

    cards_data_s2 = [
        ("Smart AI Itinerary Generation", "Utilizes Google Gemini LLM to create detailed daily travel itineraries with morning, afternoon, and evening activities tailored to user inputs.", ACCENT_BLUE),
        ("Dynamic Multi-City Routing", "Automatically partitions trip duration logically across multiple destination cities, ordering stops and recommending optimal transport.", ACCENT_INDIGO),
        ("Asynchronous BullMQ Task Queue", "Employs Redis key-value cache & BullMQ workers with multi-stage DLQ retries to eliminate API latency during heavy traffic.", ACCENT_PURPLE),
        ("Dual Map & Calendar Synchronization", "Combines an interactive 3D Canvas Globe, geographical map pinpoints, and a chronological calendar timeline for total clarity.", ACCENT_GREEN),
        ("Localized INR Currency Standardizer", "Automatically converts global currency costs ($ USD, € EUR, £ GBP, ¥ JPY) into normalized Indian Rupees (INR ₹) with locale formatting.", ACCENT_BLUE),
        ("Unified Tour Operations Dashboard", "Offers centralized analytics for tour operators and admins to monitor user trends, popular city stats, and review feedback.", ACCENT_INDIGO),
    ]

    col_w = Inches(3.8)
    row_h = Inches(2.6)
    left_positions = [Inches(0.6), Inches(4.766), Inches(8.933)]
    top_positions = [Inches(1.5), Inches(4.4)]

    for idx, (title, desc, accent) in enumerate(cards_data_s2):
        c_left = left_positions[idx % 3]
        c_top = top_positions[idx // 3]

        card = add_card(slide2, c_left, c_top, col_w, row_h)
        
        cbar = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, c_top, col_w, Inches(0.1))
        cbar.fill.solid()
        cbar.fill.fore_color.rgb = accent
        cbar.line.fill.background()

        tx = slide2.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.2), col_w - Inches(0.4), row_h - Inches(0.3))
        tf = tx.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY_TEXT
        p1.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = SECONDARY_TEXT

    # =========================================================================
    # SLIDE 3: Flow Chart / System Architecture
    # =========================================================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_background(slide3)
    add_header(slide3, "System Architecture & End-to-End Data Flow")

    workflow_steps = [
        ("1. User Request", "User inputs destinations, dates, budget & style on Next.js 15 UI", ACCENT_BLUE),
        ("2. Auth & Gateway", "Clerk verifies session; Next.js API endpoint validates payload", ACCENT_INDIGO),
        ("3. Redis & BullMQ Queue", "Checks Redis cache; if missing, enqueues job into BullMQ queue", ACCENT_PURPLE),
        ("4. Gemini AI Worker", "Worker executes Google Gemini LLM API with fallback & INR converter", ACCENT_GREEN),
        ("5. Prisma PostgreSQL", "Stores users, trips, stops, activities, budgets & reviews in Postgres", ACCENT_BLUE),
        ("6. Interactive Delivery", "Renders 3D Canvas Globe, Map view, Calendar schedule & INR breakdown", ACCENT_INDIGO),
    ]

    s3_lefts = [Inches(0.6), Inches(4.766), Inches(8.933)]
    s3_tops = [Inches(1.5), Inches(4.3)]

    for idx, (title, desc, accent) in enumerate(workflow_steps):
        c_left = s3_lefts[idx % 3]
        c_top = s3_tops[idx // 3]

        add_card(slide3, c_left, c_top, Inches(3.8), Inches(2.5))
        
        badge = slide3.shapes.add_shape(MSO_SHAPE.OVAL, c_left + Inches(0.2), c_top + Inches(0.2), Inches(0.4), Inches(0.4))
        badge.fill.solid()
        badge.fill.fore_color.rgb = accent
        badge.line.fill.background()
        
        btx = slide3.shapes.add_textbox(c_left + Inches(0.2), c_top + Inches(0.2), Inches(0.4), Inches(0.4))
        btf = btx.text_frame
        bp = btf.paragraphs[0]
        bp.text = str(idx + 1)
        bp.font.size = Pt(14)
        bp.font.bold = True
        bp.font.color.rgb = RGBColor(255, 255, 255)
        bp.alignment = PP_ALIGN.CENTER

        tx = slide3.shapes.add_textbox(c_left + Inches(0.7), c_top + Inches(0.18), Inches(2.9), Inches(2.1))
        tf = tx.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY_TEXT
        p1.space_after = Pt(6)

        p2 = tf.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(11)
        p2.font.color.rgb = SECONDARY_TEXT

    add_card(slide3, Inches(0.6), Inches(6.4), Inches(12.133), Inches(0.7), bg_color=HEADER_BG, border_color=None)
    ftx = slide3.shapes.add_textbox(Inches(0.8), Inches(6.45), Inches(11.733), Inches(0.6))
    ftf = ftx.text_frame
    fp = ftf.paragraphs[0]
    fp.text = "⚡ Key Architecture Highlight: Asynchronous execution decouples heavy LLM generation from the HTTP request/response cycle, guaranteeing instant UI response & resilient fault tolerance."
    fp.font.size = Pt(11)
    fp.font.bold = True
    fp.font.color.rgb = RGBColor(255, 255, 255)

    # =========================================================================
    # SLIDE 4: Innovation and Unique Functionality
    # =========================================================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_background(slide4)
    add_header(slide4, "Innovation & Unique Functionality")

    innovations = [
        ("Crash-Proof 3D Canvas Globe Engine", 
         "Custom-built 2D Canvas rendering engine with 3D matrix mathematics, specular shading, animated arc vectors, and pulsing markers.\n• Zero WebGL extension dependencies.\n• Prevents 'Cannot read properties of null' WebGL context crashes across low-end mobile devices and browser environments.",
         ACCENT_BLUE),
        ("Multi-Stage Asynchronous Task Queue", 
         "Production-ready Redis + BullMQ queuing pipeline featuring:\n• Automated SHA-256 request input hashing to serve cached itineraries instantly.\n• Dual-stage Dead Letter Queues (DLQ1 & DLQ2) with exponential backoff retries for resilient API fault tolerance.",
         ACCENT_INDIGO),
        ("Localized Currency Standardizer (INR ₹)", 
         "Real-time currency normalization engine:\n• Automatically detects foreign currencies ($ USD, € EUR, £ GBP, ¥ JPY, AUD, CAD) in AI outputs.\n• Converts all budget line-items into standardized Indian Rupees (INR ₹) formatted using Indian locale standards (en-IN).",
         ACCENT_GREEN),
        ("Unified Dual-View & Operations Hub", 
         "Seamlessly bridges individual traveler planning with commercial tour operations:\n• Spatial 3D Globe & Map views paired with chronological Calendar schedules.\n• Community trip discovery network for travelers & administrative oversight dashboard for tour operators.",
         ACCENT_PURPLE)
    ]

    i_positions = [
        (Inches(0.6), Inches(1.5)),
        (Inches(6.766), Inches(1.5)),
        (Inches(0.6), Inches(4.3)),
        (Inches(6.766), Inches(4.3))
    ]

    for idx, (title, text_content, accent) in enumerate(innovations):
        pos_left, pos_top = i_positions[idx]
        add_card(slide4, pos_left, pos_top, Inches(5.966), Inches(2.6))

        bar = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, pos_left, pos_top, Inches(0.12), Inches(2.6))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

        tx = slide4.shapes.add_textbox(pos_left + Inches(0.3), pos_top + Inches(0.15), Inches(5.5), Inches(2.3))
        tf = tx.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY_TEXT
        p1.space_after = Pt(4)

        for line in text_content.split('\n'):
            p = tf.add_paragraph()
            p.text = line
            if line.startswith('•'):
                p.font.size = Pt(11)
                p.font.color.rgb = SECONDARY_TEXT
            else:
                p.font.size = Pt(12)
                p.font.color.rgb = PRIMARY_TEXT
                p.font.bold = True
            p.space_after = Pt(2)

    # =========================================================================
    # SLIDE 5: Technical Details
    # =========================================================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_background(slide5)
    add_header(slide5, "Technical Details — Stack, Deployment & Cost")

    add_card(slide5, Inches(0.6), Inches(1.5), Inches(3.8), Inches(5.5))
    tx_t1 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(3.4), Inches(5.1))
    tf_t1 = tx_t1.text_frame
    tf_t1.word_wrap = True

    p = tf_t1.paragraphs[0]
    p.text = "FRAMEWORKS & TECH"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    tech_items = [
        ("Frontend Framework", "Next.js 15 (App Router), React 19, TypeScript"),
        ("Styling & UI", "Tailwind CSS, Shadcn UI, Framer Motion, Lucide Icons"),
        ("Task Queue & Cache", "Redis (ioredis), BullMQ Worker Queues"),
        ("AI Generation", "Google Generative AI (Gemini LLM API)"),
        ("Database & ORM", "PostgreSQL, Prisma ORM"),
        ("Authentication", "Clerk Authentication (Social & Email Login)"),
        ("Visualization Engine", "Custom HTML5 Canvas 2D/3D Engine")
    ]

    for label, val in tech_items:
        p = tf_t1.add_paragraph()
        p.text = label + ":"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_TEXT

        p_sub = tf_t1.add_paragraph()
        p_sub.text = val
        p_sub.font.size = Pt(10)
        p_sub.font.color.rgb = SECONDARY_TEXT
        p_sub.space_after = Pt(4)

    add_card(slide5, Inches(4.766), Inches(1.5), Inches(3.8), Inches(5.5))
    tx_t2 = slide5.shapes.add_textbox(Inches(4.966), Inches(1.7), Inches(3.4), Inches(5.1))
    tf_t2 = tx_t2.text_frame
    tf_t2.word_wrap = True

    p = tf_t2.paragraphs[0]
    p.text = "DEPLOYMENT SETUP"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO

    deploy_items = [
        ("Web Application Host", "Vercel Cloud Platform\n(Serverless & Edge Next.js Functions)"),
        ("Redis Cache & Queues", "Upstash Serverless Redis\n(Fully managed, low latency key-value store)"),
        ("Database Hosting", "Neon PostgreSQL Serverless\n(Auto-scaling relational cloud database)"),
        ("AI Service Deployment", "Google Cloud Gemini API Gateway\n(High-throughput LLM endpoints)"),
        ("CI/CD Pipeline", "GitHub Actions & Vercel Automated Preview Builds")
    ]

    for label, val in deploy_items:
        p = tf_t2.add_paragraph()
        p.text = label + ":"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_TEXT

        p_sub = tf_t2.add_paragraph()
        p_sub.text = val
        p_sub.font.size = Pt(10)
        p_sub.font.color.rgb = SECONDARY_TEXT
        p_sub.space_after = Pt(4)

    add_card(slide5, Inches(8.933), Inches(1.5), Inches(3.8), Inches(5.5))
    tx_t3 = slide5.shapes.add_textbox(Inches(9.133), Inches(1.7), Inches(3.4), Inches(5.1))
    tf_t3 = tx_t3.text_frame
    tf_t3.word_wrap = True

    p = tf_t3.paragraphs[0]
    p.text = "ESTIMATED OPERATIONAL COST"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    costs = [
        ("Vercel Hosting", "$0 - $20 / month"),
        ("Upstash Redis", "$0 - $10 / month"),
        ("Neon PostgreSQL", "$0 - $25 / month"),
        ("Gemini AI API", "Free Tier / Pay-as-you-go (~$0.0001 per call)"),
        ("Clerk Authentication", "Free Tier (up to 10,000 MAU)"),
    ]

    for service, cost in costs:
        p = tf_t3.add_paragraph()
        p.text = service
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_TEXT

        p_sub = tf_t3.add_paragraph()
        p_sub.text = cost
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = ACCENT_GREEN
        p_sub.font.bold = True
        p_sub.space_after = Pt(6)

    total_box = tf_t3.add_paragraph()
    total_box.text = "TOTAL ESTIMATED COST:"
    total_box.font.size = Pt(12)
    total_box.font.bold = True
    total_box.font.color.rgb = PRIMARY_TEXT

    total_val = tf_t3.add_paragraph()
    total_val.text = "~$0 – $55 / Month"
    total_val.font.size = Pt(20)
    total_val.font.bold = True
    total_val.font.color.rgb = ACCENT_BLUE

    # =========================================================================
    # SLIDE 6: Existing Solutions and Comparison
    # =========================================================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_background(slide6)
    add_header(slide6, "Existing Solutions and Comparison Matrix")

    rows, cols = 7, 6
    left, top, width, height = Inches(0.6), Inches(1.5), Inches(12.133), Inches(5.5)
    table_shape = slide6.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    table.columns[0].width = Inches(2.633)
    table.columns[1].width = Inches(2.1)
    table.columns[2].width = Inches(1.8)
    table.columns[3].width = Inches(1.8)
    table.columns[4].width = Inches(1.9)
    table.columns[5].width = Inches(1.9)

    headers = [
        "Feature / Dimension", 
        "Tripzzy (Roamly)", 
        "TripIt", 
        "Wanderlog", 
        "Generic ChatGPT", 
        "Tour Operators"
    ]

    matrix_data = [
        ["Real-time AI Itinerary Synthesis", "✅ Instant & Dynamic", "❌ Manual Forwarding", "⚠️ Limited AI", "⚠️ Text Only", "❌ Slow / Fixed"],
        ["Asynchronous Queue (BullMQ)", "✅ Built-in Queue", "❌ None", "❌ None", "❌ None", "❌ None"],
        ["Crash-Proof 3D Globe Visualizer", "✅ Interactive 3D Canvas", "❌ Static Maps", "⚠️ Basic 2D Pin", "❌ Text Only", "❌ Printed Papers"],
        ["Localized Currency (INR ₹)", "✅ Automated Conversion", "❌ USD / Local Only", "⚠️ Mixed Currency", "⚠️ Manual Prompting", "⚠️ Fixed Quotes"],
        ["Multi-City Route Auto-Optimizer", "✅ Smart Logical Splits", "❌ Manual Layout", "⚠️ Semi-Automated", "⚠️ Inconsistent", "⚠️ Pre-packaged"],
        ["Integrated Admin & Community", "✅ Analytics + Sharing", "❌ Personal Only", "⚠️ Basic Sharing", "❌ None", "❌ Closed Systems"]
    ]

    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEADER_BG if col_idx != 1 else ACCENT_BLUE
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)
        p.alignment = PP_ALIGN.CENTER

    for row_idx, row_content in enumerate(matrix_data):
        for col_idx, text in enumerate(row_content):
            cell = table.cell(row_idx + 1, col_idx)
            cell.fill.solid()
            if col_idx == 0:
                cell.fill.fore_color.rgb = RGBColor(241, 245, 249)
            elif col_idx == 1:
                cell.fill.fore_color.rgb = RGBColor(239, 246, 255)
            else:
                cell.fill.fore_color.rgb = RGBColor(255, 255, 255) if row_idx % 2 == 0 else RGBColor(248, 250, 252)

            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.text = text
            p.font.size = Pt(11)
            p.font.name = "Calibri"
            if col_idx == 0:
                p.font.bold = True
                p.font.color.rgb = PRIMARY_TEXT
                p.alignment = PP_ALIGN.LEFT
            elif col_idx == 1:
                p.font.bold = True
                p.font.color.rgb = ACCENT_BLUE
                p.alignment = PP_ALIGN.CENTER
            else:
                p.font.color.rgb = SECONDARY_TEXT
                p.alignment = PP_ALIGN.CENTER

    output_filename = "Tripzzy_PS07_Ghayal_Ghoda_Presentation.pptx"
    output_path = os.path.join(os.getcwd(), output_filename)
    prs.save(output_path)
    print(f"Presentation saved successfully at: {output_path}")

    # Copy to artifacts directory if available
    artifact_dir = r"C:\Users\Atharva\.gemini\antigravity-ide\brain\00846d85-d373-421f-bf87-f6307f487954"
    if os.path.exists(artifact_dir):
        import shutil
        shutil.copy(output_path, os.path.join(artifact_dir, output_filename))
        print(f"Presentation copied to artifacts directory: {os.path.join(artifact_dir, output_filename)}")

if __name__ == "__main__":
    build_presentation()
